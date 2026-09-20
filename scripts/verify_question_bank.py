"""Independent source-to-JSON checks before assessment implementation."""
import hashlib
import json
from pathlib import Path
import unittest
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BANK = json.loads((ROOT/'web/data/question-bank.json').read_text())
SOURCE = ROOT/'source'/BANK['source_file']
WORKBOOK = openpyxl.load_workbook(SOURCE, data_only=False)

class QuestionBankTests(unittest.TestCase):
    def test_source_checksum(self):
        self.assertEqual(hashlib.sha256(SOURCE.read_bytes()).hexdigest(), BANK['source_sha256'])

    def test_all_items_match_source_verbatim(self):
        rows = list(WORKBOOK['题库'].iter_rows(min_row=6, values_only=True))
        self.assertEqual(len(rows), len(BANK['items']))
        for index, (row, item) in enumerate(zip(rows, BANK['items']), 1):
            self.assertEqual((row[0], row[1], row[2], row[3], row[8]),
                             (item['item_id'], item['dimension'], item['subdimension'], item['question_text'], item['dimension_key']))
            self.assertEqual(item['order'], index)
            self.assertEqual(item['scoring_direction'], row[4] or None)
            self.assertEqual(item['reverse_scoring'], row[5] == '是')

    def test_count_and_endpoints(self):
        self.assertEqual(len(BANK['items']), 72)
        self.assertEqual(BANK['items'][0]['item_id'], 'Q01')
        self.assertEqual(BANK['items'][-1]['item_id'], 'Q72')
        self.assertEqual(BANK['items'][0]['question_text'], WORKBOOK['题库']['D6'].value)
        self.assertEqual(BANK['items'][-1]['question_text'], WORKBOOK['题库']['D77'].value)

    def test_ids(self):
        ids = [i['item_id'] for i in BANK['items']]
        self.assertEqual(len(set(ids)), len(ids))
        self.assertEqual(ids, [f'Q{i:02}' for i in range(1, 73)])

    def test_student_options_only(self):
        rows = list(WORKBOOK['作答选项'].iter_rows(min_row=6, values_only=True))
        self.assertEqual(BANK['options'], [{'value': r[0], 'label': r[1]} for r in rows])

    def test_scoring_is_not_invented(self):
        self.assertEqual(BANK['scoring_policy']['status'], 'awaiting_confirmation')
        self.assertIsNone(BANK['scoring_policy']['aggregation'])
        self.assertTrue(all(not i['reverse_scoring'] for i in BANK['items']))

if __name__ == '__main__':
    unittest.main(verbosity=2)
